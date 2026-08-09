import json
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from docx import Document
from llama_index.core.schema import NodeWithScore, TextNode

from functions import rag
from functions.rag import load_document_file, load_documents


def test_agent_planner_requests_schema_from_both_providers(monkeypatch):
    requests = []

    class FakeCompletions:
        def create(self, **request):
            requests.append(request)
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(
                            content='{"query":"standalone","intent":"answer"}'
                        )
                    )
                ],
                usage=None,
            )

    monkeypatch.setattr(
        rag,
        "OpenAI",
        lambda **_kwargs: SimpleNamespace(
            chat=SimpleNamespace(completions=FakeCompletions())
        ),
    )
    monkeypatch.setenv("GROQ_API_KEY", "test-key")

    for provider in ("ollama", "groq"):
        monkeypatch.setenv("AGENTIC_CRAG_LLM_PROVIDER", provider)
        text, _ = rag.call_model("plan this", node="planner")
        assert text == '{"query":"standalone","intent":"answer"}'

    for request in requests:
        response_format = request["response_format"]
        assert response_format["type"] == "json_schema"
        assert response_format["json_schema"]["strict"] is True
        assert response_format["json_schema"]["schema"] == {
            "type": "object",
            "properties": {
                "query": {
                    "type": "string",
                    "minLength": 1,
                    "maxLength": 500,
                },
                "intent": {
                    "type": "string",
                    "enum": ["answer", "overview"],
                },
            },
            "required": ["query", "intent"],
            "additionalProperties": False,
        }


def test_every_control_schema_is_requested_from_both_providers(monkeypatch):
    requests = []

    class FakeCompletions:
        def create(self, **request):
            requests.append(request)
            return SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content="{}"))],
                usage=None,
            )

    monkeypatch.setattr(
        rag,
        "OpenAI",
        lambda **_kwargs: SimpleNamespace(
            chat=SimpleNamespace(completions=FakeCompletions())
        ),
    )
    monkeypatch.setenv("GROQ_API_KEY", "test-key")

    expected = []
    for provider in ("ollama", "groq"):
        monkeypatch.setenv("AGENTIC_CRAG_LLM_PROVIDER", provider)
        for node, schema in rag.STRUCTURED_OUTPUT_SCHEMAS.items():
            rag.call_model("control", node=node)
            expected.append((node, schema))

    assert len(requests) == len(expected)
    for request, (node, schema) in zip(requests, expected, strict=True):
        value = request["response_format"]["json_schema"]
        assert value["name"] == f"agentic_crag_{node}"
        assert value["strict"] is True
        assert value["schema"] == schema


def test_load_docx(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Agentic CRAG project overview")
    doc.save(path)

    docs = load_document_file(path)

    assert docs == [
        {
            "filename": "sample.docx",
            "text": "Document filename: sample.docx\n\nAgentic CRAG project overview",
            "type": "docx",
        }
    ]


def test_load_xlsx(tmp_path: Path):
    path = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Amount"])
    ws.append(["Hosting", 25])
    wb.save(path)

    docs = load_document_file(path)

    assert len(docs) == 1
    assert docs[0]["filename"] == "sample.xlsx-Budget"
    assert docs[0]["type"] == "xlsx"
    assert "Item: Hosting | Amount: 25" in docs[0]["text"]


def test_load_image_uses_ocr(tmp_path: Path, monkeypatch):
    path = tmp_path / "label.png"
    path.write_bytes(b"image bytes")
    monkeypatch.setattr(rag, "_ocr_image_text", lambda image: "Asset ID ZX-9001")

    docs = load_document_file(path)

    assert docs == [
        {
            "filename": "label.png",
            "text": "Document filename: label.png\n\nAsset ID ZX-9001",
            "type": "image",
        }
    ]


def test_scanned_pdf_pages_fall_back_to_ocr(tmp_path: Path, monkeypatch):
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"pdf bytes")
    pages = [
        SimpleNamespace(
            extract_text=lambda: "Embedded text is already long enough to keep."
        ),
        SimpleNamespace(extract_text=lambda: ""),
    ]
    monkeypatch.setattr(rag, "PdfReader", lambda _path: SimpleNamespace(pages=pages))
    monkeypatch.setattr(
        rag,
        "_ocr_pdf_pages",
        lambda _path, page_numbers: {page_numbers[0]: "Scanned asset ZX-9001"},
    )

    docs = load_document_file(path)

    assert "Page 1:\nEmbedded text is already long enough to keep." in docs[0]["text"]
    assert "Page 2:\nScanned asset ZX-9001" in docs[0]["text"]


def test_ocr_engine_is_loaded_once(monkeypatch):
    engines = []
    monkeypatch.setattr(rag, "_ocr_engine", None)
    monkeypatch.setattr(
        rag, "_build_ocr_engine", lambda: engines.append(object()) or engines[-1]
    )

    assert rag._get_ocr_engine() is rag._get_ocr_engine()
    assert len(engines) == 1


def test_unsupported_file_is_reported(tmp_path: Path):
    (tmp_path / "notes.txt").write_text("ignore me", encoding="utf-8")

    docs, warnings = load_documents(tmp_path)

    assert docs == []
    assert warnings == ["Skipped unsupported file: notes.txt"]


def test_remove_index_removes_registry_entry_and_index_dir(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("AGENTIC_CRAG_STORAGE_DIR", str(tmp_path))
    index_dir = rag.get_indexes_dir() / "demo"
    index_dir.mkdir(parents=True)
    registry_file = tmp_path / "registry.json"
    registry_file.write_text(
        json.dumps({"demo": {"documents": []}, "keep": {"documents": []}}),
        encoding="utf-8",
    )

    assert rag.remove_index("demo") is True
    assert not index_dir.exists()
    assert json.loads(registry_file.read_text(encoding="utf-8")) == {
        "keep": {"documents": []}
    }


def test_embedding_setup_reuses_one_model(monkeypatch):
    created = []
    settings = SimpleNamespace(embed_model=None)

    def embedding(model_name):
        model = object()
        created.append((model_name, model))
        return model

    monkeypatch.setattr(rag, "_embedding_model", None)
    monkeypatch.setattr(rag, "Settings", settings)
    monkeypatch.setattr(rag, "HuggingFaceEmbedding", embedding)

    rag.setup_embeddings()
    first = settings.embed_model
    rag.setup_embeddings()

    assert created == [("BAAI/bge-small-en-v1.5", first)]
    assert settings.embed_model is first


def test_load_index_reuses_cached_instance(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("AGENTIC_CRAG_STORAGE_DIR", str(tmp_path))
    index_dir = rag.get_indexes_dir() / "demo"
    index_dir.mkdir(parents=True)
    loaded = []
    expected = object()

    monkeypatch.setattr(rag, "_loaded_indexes", {})
    monkeypatch.setattr(rag, "setup_embeddings", lambda: None)
    monkeypatch.setattr(
        rag.StorageContext,
        "from_defaults",
        lambda persist_dir: SimpleNamespace(persist_dir=persist_dir),
    )
    monkeypatch.setattr(
        rag,
        "load_index_from_storage",
        lambda storage: loaded.append(storage.persist_dir) or expected,
    )

    assert rag.load_index("demo") is expected
    assert rag.load_index("demo") is expected
    assert loaded == [str(index_dir)]


def test_answer_entrypoint_uses_shared_retrieval_and_model_call(monkeypatch):
    calls = []
    monkeypatch.setattr(
        rag,
        "retrieve_sources",
        lambda _index, query, top_k: calls.append(("retrieve", query, top_k)) or [],
    )
    monkeypatch.setattr(
        rag,
        "call_model",
        lambda _prompt, *, node, timeout: (
            calls.append(("complete", node, timeout))
            or ("The answer is not present in the provided documents.", 1)
        ),
    )

    result = rag.ask_index_with_sources(object(), "Unknown fact?")

    assert calls == [
        ("retrieve", "Unknown fact?", rag.RERANK_TOP_N),
        ("complete", "synthesis", 30.0),
    ]
    assert result["answer"] == "The answer is not present in the provided documents."


def test_direct_rag_uses_shared_conversational_overview(monkeypatch):
    source = {
        "filename": "incident.png",
        "type": "image",
        "score": 0.99,
        "text": (
            "INCIDENT REPORT Incident ID: OCR-417 Affected service: "
            "document-ingestion Root cause: OCR batches exceeded the worker memory "
            "limit. Resolution: Limit each OCR batch to eight pages. Status: Resolved"
        ),
    }
    queries = []
    monkeypatch.setattr(
        rag,
        "retrieve_sources",
        lambda _index, query, top_k: queries.append((query, top_k)) or [source],
    )
    called_nodes = []

    def call_model(_prompt, node, **_kwargs):
        called_nodes.append(node)
        responses = {
            "planner": '{"query":"OCR-417 incident report overview","intent":"overview"}',
            "overview": json.dumps(
                {
                    "answer": (
                        "This incident report covers OCR-417. The document-ingestion "
                        "service was affected because OCR batches exceeded the worker "
                        "memory limit. The issue was resolved by limiting each OCR batch "
                        "to eight pages. The status is resolved."
                    )
                }
            ),
        }
        return responses[node], 10

    monkeypatch.setattr(rag, "call_model", call_model)

    result = rag.ask_index_with_sources(
        object(),
        "What is this about?",
        history=[
            {"role": "user", "content": "Tell me about incident OCR-417."},
            {
                "role": "assistant",
                "content": "OCR-417 affected document-ingestion [incident.png].",
                "source_filenames": ["incident.png"],
            },
        ],
    )

    assert queries == [
        ("OCR-417 incident report overview", rag.RERANK_TOP_N)
    ]
    assert called_nodes == ["planner", "overview"]
    assert result["answer"] == (
        "This incident report covers OCR-417. The document-ingestion service was "
        "affected because OCR batches exceeded the worker memory limit. The issue was "
        "resolved by limiting each OCR batch to eight pages. The status is resolved "
        "[incident.png]."
    )


def test_direct_rag_uses_shared_structured_multi_field_answer(monkeypatch):
    source = {
        "filename": "incident.png",
        "type": "image",
        "score": 0.99,
        "text": (
            "INCIDENT REPORT Incident ID: OCR-417 Affected service: "
            "document-ingestion Root cause: OCR batches exceeded the worker memory "
            "limit. Resolution: Limit each OCR batch to eight pages. Status: Resolved"
        ),
    }
    monkeypatch.setattr(rag, "retrieve_sources", lambda *_args, **_kwargs: [source])
    monkeypatch.setattr(
        rag,
        "call_model",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            AssertionError("structured fields should not call the model")
        ),
    )

    result = rag.ask_index_with_sources(
        object(),
        (
            "For incident OCR-417, which service was affected, what caused it, "
            "and how was it resolved?"
        ),
    )

    assert result["answer"] == (
        "For incident OCR-417, the report says [incident.png]:\n"
        "- Affected service: document-ingestion\n"
        "- What caused it: OCR batches exceeded the worker memory limit.\n"
        "- How it was resolved: Limit each OCR batch to eight pages."
    )


def test_supporting_sources_follow_citation_order_and_hide_retrieval_scores():
    sources = [
        {
            "filename": "runbook.docx",
            "type": "docx",
            "score": 0.91,
            "text": "Nginx listens publicly on port 80.",
        },
        {
            "filename": "release.docx",
            "type": "docx",
            "score": 0.04,
            "text": "Windows uses .agentic_crag_data and EC2 uses /opt/data.",
        },
        {
            "filename": "noise.docx",
            "type": "docx",
            "score": 0.99,
            "text": "This passage is retrieved but does not support the answer.",
        },
    ]

    grouped = rag.supporting_source_groups(
        (
            "Windows uses .agentic_crag_data and EC2 uses /opt/data "
            "[release.docx]. Nginx listens publicly on port 80 [runbook.docx]."
        ),
        sources,
    )

    assert [source["filename"] for source in grouped] == [
        "release.docx",
        "runbook.docx",
    ]
    assert all("score" not in source for source in grouped)
    assert grouped[0]["passages"] == [
        {"text": "Windows uses .agentic_crag_data and EC2 uses /opt/data."}
    ]


def test_supporting_sources_extract_the_exact_cited_passage():
    grouped = rag.supporting_source_groups(
        "The configured value is 42 [requirements.docx].",
        [
            {
                "filename": "requirements.docx",
                "type": "docx",
                "score": 0.9,
                "text": (
                    "Unrelated introductory material. "
                    "The configured value is 42. "
                    "Unrelated closing material."
                ),
            }
        ],
    )

    assert grouped[0]["passages"] == [
        {"text": "The configured value is 42."}
    ]


def test_supporting_sources_reject_a_cited_document_that_does_not_support_claim():
    grouped = rag.supporting_source_groups(
        "The configured value is 42 [noise.docx].",
        [
            {
                "filename": "noise.docx",
                "type": "docx",
                "score": 0.99,
                "text": "This document only discusses deployment scheduling.",
            }
        ],
    )

    assert grouped == []


def test_supporting_sources_validate_the_claim_attached_to_each_citation():
    grouped = rag.supporting_source_groups(
        (
            "The configured value is 42 [requirements.docx]. "
            "Deployment happens on Tuesday [noise.docx]."
        ),
        [
            {
                "filename": "requirements.docx",
                "type": "docx",
                "text": "The configured value is 42.",
            },
            {
                "filename": "noise.docx",
                "type": "docx",
                "text": "An archive also records the configured value as 42.",
            },
        ],
    )

    assert [source["filename"] for source in grouped] == ["requirements.docx"]


def test_direct_rag_returns_grouped_supporting_documents(monkeypatch):
    incident = {
        "filename": "incident.png",
        "type": "image",
        "score": 0.2,
        "text": (
            "Incident ID: OCR-417 Affected service: document-ingestion "
            "Root cause: memory pressure."
        ),
    }
    noise = {
        "filename": "noise.docx",
        "type": "docx",
        "score": 0.99,
        "text": "Unrelated deployment notes.",
    }
    monkeypatch.setattr(
        rag, "retrieve_sources", lambda *_args, **_kwargs: [noise, incident]
    )
    monkeypatch.setattr(
        rag,
        "call_model",
        lambda *_args, **_kwargs: (
            "OCR-417 affected document-ingestion [incident.png].",
            10,
        ),
    )

    result = rag.ask_index_with_sources(object(), "Which service did OCR-417 affect?")

    assert [source["filename"] for source in result["sources"]] == ["incident.png"]
    assert result["sources"][0]["passages"] == [
        {"text": "Incident ID: OCR-417 Affected service: document-ingestion"}
    ]
    assert "score" not in result["sources"][0]


def test_spreadsheet_lookup_matches_requested_row_and_column():
    source = {
        "filename": "budget.xlsx-Quarterly Budget",
        "type": "xlsx",
        "text": (
            "Quarter: Q3 | Category: EC2 hosting | Budget_USD: 30 | Actual_USD: 0\n"
            "Quarter: Q3 | Category: Groq API | Budget_USD: 75 | Actual_USD: 0"
        ),
    }

    assert (
        rag.answer_spreadsheet_lookup(
            "What is the projected Q3 Groq API budget?", [source]
        )
        == "$75 [budget.xlsx-Quarterly Budget]"
    )


def test_spreadsheet_lookup_ignores_notes_when_matching_a_row():
    source = {
        "filename": "budget.xlsx-Quarterly Budget",
        "type": "xlsx",
        "text": (
            "Quarter: Q1 | Category: Groq API | Budget_USD: 40 | Actual_USD: 28 | "
            "Notes: Used for final answers\n"
            "Quarter: Q3 | Category: Groq API | Budget_USD: 75 | Actual_USD: 0 | "
            "Notes: Projected token spend"
        ),
    }

    assert (
        rag.answer_spreadsheet_lookup(
            "What was the Q1 Groq API budget and actual spend?", [source]
        )
        == "Budget: $40; Actual: $28 [budget.xlsx-Quarterly Budget]"
    )


def test_spreadsheet_lookup_uses_requested_latency_field():
    source = {
        "filename": "benchmarks.xlsx-Indexing Benchmarks",
        "type": "xlsx",
        "text": (
            "Document_Set: demo-corpus | Index_Time_Seconds: 42.8 | Retrieval_Time_ms: 240\n"
            "Document_Set: spreadsheet-heavy | Index_Time_Seconds: 58.1 | Retrieval_Time_ms: 295"
        ),
    }

    assert (
        rag.answer_spreadsheet_lookup(
            "What retrieval time was recorded for the spreadsheet-heavy benchmark?",
            [source],
        )
        == "295 ms [benchmarks.xlsx-Indexing Benchmarks]"
    )


def test_formatted_spreadsheet_source_preserves_rows_and_full_sheet():
    class FakeNode:
        metadata = {"filename": "sheet.xlsx-Data", "type": "xlsx"}

        @staticmethod
        def get_content(metadata_mode="none"):
            assert metadata_mode == "none"
            return "A: first | B: 1\nA: second | B: 2\n" + ("x" * 3000)

    text = rag.format_source_nodes([FakeNode()])[0]["text"]

    assert "A: first | B: 1\nA: second | B: 2" in text
    assert len(text) > rag.MAX_SOURCE_TEXT_CHARS


def test_agent_retrieval_fetches_wide_candidates_then_reranks(monkeypatch):
    calls = {}
    candidate = NodeWithScore(node=TextNode(text="semantic candidate"), score=0.8)

    class Retriever:
        @staticmethod
        def retrieve(query):
            calls["retrieve_query"] = query
            return [candidate]

    class Index:
        docstore = SimpleNamespace(docs={})

        @staticmethod
        def as_retriever(similarity_top_k):
            calls["similarity_top_k"] = similarity_top_k
            return Retriever()

    class Reranker:
        @staticmethod
        def postprocess_nodes(nodes, query_str):
            calls["rerank"] = (nodes, query_str)
            return []

    def get_reranker(top_n):
        calls["top_n"] = top_n
        return Reranker()

    monkeypatch.setattr(rag, "get_reranker", get_reranker)

    assert rag.retrieve_sources(Index(), "FR-006", top_k=5) == []
    assert calls["similarity_top_k"] == rag.RETRIEVE_TOP_K
    assert calls["retrieve_query"] == "FR-006"
    assert calls["top_n"] == rag.RETRIEVE_TOP_K
    fused, rerank_query = calls["rerank"]
    assert rerank_query == "FR-006"
    assert [item.node_id for item in fused] == [candidate.node_id]


def test_hybrid_retrieval_promotes_exact_identifier(monkeypatch):
    semantic = TextNode(
        text="A general deployment configuration overview.",
        metadata={"filename": "overview.docx", "type": "docx"},
    )
    exact = TextNode(
        text="Asset ZX-9001 belongs to the cobalt configuration.",
        metadata={"filename": "assets.docx", "type": "docx"},
    )

    class Retriever:
        @staticmethod
        def retrieve(_query):
            return [NodeWithScore(node=semantic, score=0.99)]

    class Index:
        docstore = SimpleNamespace(
            docs={semantic.node_id: semantic, exact.node_id: exact}
        )

        @staticmethod
        def as_retriever(similarity_top_k):
            assert similarity_top_k == rag.RETRIEVE_TOP_K
            return Retriever()

    class Reranker:
        @staticmethod
        def postprocess_nodes(nodes, query_str):
            assert query_str == "Which configuration belongs to ZX-9001?"
            return nodes

    monkeypatch.setattr(rag, "_keyword_indexes", {})
    monkeypatch.setattr(rag, "get_reranker", lambda top_n: Reranker())

    sources = rag.retrieve_sources(
        Index(), "Which configuration belongs to ZX-9001?", top_k=2
    )

    assert [source["filename"] for source in sources] == [
        "assets.docx",
        "overview.docx",
    ]


def test_spreadsheet_excerpt_keeps_rows_relevant_to_query():
    text = "\n".join(
        [
            "Document filename: budget.xlsx",
            "Quarter: Q1 | Category: Hosting | Budget_USD: 18",
            "Quarter: Q2 | Category: Hosting | Budget_USD: 22",
            "Quarter: Q3 | Category: Hosting | Budget_USD: 30",
            "Quarter: Q3 | Category: Groq API | Budget_USD: 75",
            "Quarter: Q4 | Category: Database | Budget_USD: 120",
        ]
    )

    excerpt = rag._spreadsheet_excerpt(text, "Q3 Groq API budget", max_rows=2)

    assert "Quarter: Q3 | Category: Groq API | Budget_USD: 75" in excerpt
    assert excerpt.count("Quarter:") == 2


def test_spreadsheet_excerpt_keeps_all_rows_for_analytical_query():
    text = "\n".join(
        [
            "Document filename: benchmarks.xlsx",
            "Document_Set: one | Retrieval_Time_ms: 10",
            "Document_Set: two | Retrieval_Time_ms: 20",
            "Document_Set: three | Retrieval_Time_ms: 30",
        ]
    )

    excerpt = rag._spreadsheet_excerpt(
        text, "Which benchmark has the slowest retrieval time?", max_rows=1
    )

    assert excerpt == text
